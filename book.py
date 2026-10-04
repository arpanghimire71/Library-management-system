import StudentData as sd
from openpyxl import load_workbook as lw
def book():
    sd.master()
    if sd.status!="Logged in":
        print("Access denied")
        exit()
    if sd.status=="Logged in":
        wb=lw("C:/Users/Lenovo/OneDrive/Desktop/Library management software/Book information.xlsx")
        sheet=wb["Sheet1"]
        print("Book information:")
        print("Book id\t\t Book name\t\tquantity in library")
   
    if sd.status=="Logged in":
        for row in sheet.iter_rows(min_row=2,values_only=True):
            book_name,quantity,book_id=row
            print(f"{book_id}\t\t{book_name}\t\t\t{quantity}")
            quantity=int(quantity)
def borrow():
    if sd.status=="Logged in":
        book_code=input("Enter the book code you want to borrow:")
        wb=lw("C:/Users/Lenovo/OneDrive/Desktop/Library management software/Book information.xlsx")
        sheet=wb["Sheet1"]
        for row in sheet.iter_rows(min_row=2):
            book_name,quantity,book_id=row
            if str(book_id.value)==str(book_code):
                quantity1=int(quantity.value)
                if quantity1>0:
                    quantity1=quantity1-1
                    quantity.value=quantity1
                    wb.save("C:/Users/Lenovo/OneDrive/Desktop/Library management software/Book information.xlsx")
                    print(f"You have borrowed {book_name.value}. Please return it within 15 days.")
                else:
                    print(f"Sorry, {book_name.value} is not available in the library.")

def return_book():
    if sd.status=="Logged in":
        book_code=input("Enter the book code you want to return:")
        wb=lw("C:/Users/Lenovo/OneDrive/Desktop/Library management software/Book information.xlsx")
        sheet=wb["Sheet1"]
        for row in sheet.iter_rows(min_row=2):
            book_name,quantity,book_id=row
            if str(book_id.value)==str(book_code):
                quantity1=int(quantity.value)
                quantity1=quantity1+1
                quantity.value=quantity1
                wb.save("C:/Users/Lenovo/OneDrive/Desktop/Library management software/Book information.xlsx")
                print(f"You have returned {book_name.value}. Thank you for returning it on time.")

if __name__=="__main__":
    book()

requirements=input("Enter 1 to borrow\n Enter 2 to return\nEnter 3 to exit\n")
if requirements=="1":
    if sd.status=="Logged in":
        borrow()
    else:
        exit()   
if requirements=="2":
    if sd.status=="Logged in":
        return_book()
    else:
        exit()   
if requirements=="3":
    print("Exitted successfully")
    exit()
